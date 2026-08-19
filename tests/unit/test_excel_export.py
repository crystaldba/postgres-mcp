"""Tests for Excel export functionality (format_to_excel and execute_sql_xlsx tool)."""

import os
import tempfile
from typing import Any
from unittest.mock import AsyncMock
from unittest.mock import MagicMock
from unittest.mock import patch

import pytest
import pytest_asyncio
from mcp import types

import postgres_mcp.server as server
from postgres_mcp.formatter import format_to_excel
from postgres_mcp.server import AccessMode
from postgres_mcp.sql import SafeSqlDriver


class MockRowResult:
    """Mock row matching SqlDriver.RowResult interface."""

    def __init__(self, cells: dict[str, Any]):
        self.cells = cells


def response_text(result: server.ResponseType) -> str:
    """Return text from a single-text-content tool response."""
    assert len(result) == 1
    content = result[0]
    assert isinstance(content, types.TextContent)
    return content.text


# ---------------------------------------------------------------------------
# format_to_excel tests
# ---------------------------------------------------------------------------


def test_format_to_excel_creates_file():
    """format_to_excel creates a valid xlsx file with data."""
    rows = [{"id": 1, "name": "Alice"}, {"id": 2, "name": "Bob"}]
    columns = ["id", "name"]

    with tempfile.TemporaryDirectory() as tmpdir:
        path = format_to_excel(rows, columns, output_dir=tmpdir)
        assert os.path.exists(path)
        assert path.endswith(".xlsx")
        assert tmpdir in path


def test_format_to_excel_custom_output_dir():
    """format_to_excel uses the provided output directory."""
    rows = [{"x": 10}]
    columns = ["x"]

    with tempfile.TemporaryDirectory() as tmpdir:
        path = format_to_excel(rows, columns, output_dir=tmpdir)
        assert os.path.dirname(path) == tmpdir


def test_format_to_excel_default_output_dir():
    """format_to_excel defaults to system temp / postgres-mcp-results."""
    rows = [{"a": 1}]
    columns = ["a"]

    path = format_to_excel(rows, columns)
    expected_dir = os.path.join(tempfile.gettempdir(), "postgres-mcp-results")
    assert path.startswith(expected_dir)
    assert os.path.exists(path)
    os.unlink(path)


def test_format_to_excel_column_width_capped():
    """format_to_excel caps column width at 50."""
    rows = [{"short": "hi", "long": "a" * 100}]
    columns = ["short", "long"]

    with tempfile.TemporaryDirectory() as tmpdir:
        path = format_to_excel(rows, columns, output_dir=tmpdir)

        from openpyxl import load_workbook

        wb = load_workbook(path)
        ws = wb.active
        assert ws is not None
        assert ws["A1"].value == "short"
        assert ws["B1"].value == "long"
        assert ws["A2"].value == "hi"
        assert ws.column_dimensions["B"].width <= 50


def test_format_to_excel_handles_none_values():
    """format_to_excel handles None values in row data."""
    rows = [{"id": 1, "name": None}, {"id": None, "name": "test"}]
    columns = ["id", "name"]

    with tempfile.TemporaryDirectory() as tmpdir:
        path = format_to_excel(rows, columns, output_dir=tmpdir)

        from openpyxl import load_workbook

        wb = load_workbook(path)
        ws = wb.active
        assert ws is not None
        assert ws["B2"].value is None
        assert ws["A3"].value is None


def test_format_to_excel_empty_rows():
    """format_to_excel handles empty row list (creates header-only file)."""
    rows = []
    columns = ["id", "name"]

    with tempfile.TemporaryDirectory() as tmpdir:
        path = format_to_excel(rows, columns, output_dir=tmpdir)

        from openpyxl import load_workbook

        wb = load_workbook(path)
        ws = wb.active
        assert ws is not None
        assert ws["A1"].value == "id"
        assert ws.max_row == 1  # header only


def test_format_to_excel_serializes_complex_types():
    """format_to_excel serializes dict and list values to JSON strings."""
    rows = [
        {"id": 1, "json_col": {"key": "value"}, "arr_col": [1, 2, 3]},
    ]
    columns = ["id", "json_col", "arr_col"]

    with tempfile.TemporaryDirectory() as tmpdir:
        path = format_to_excel(rows, columns, output_dir=tmpdir)

        from openpyxl import load_workbook

        wb = load_workbook(path)
        ws = wb.active
        assert ws is not None
        # dict -> JSON string
        assert ws["B2"].value == '{"key": "value"}'
        # list -> JSON string
        assert ws["C2"].value == "[1, 2, 3]"


def test_format_to_excel_filename_is_unique():
    """format_to_excel generates unique filenames to avoid concurrent overwrites."""
    rows = [{"a": 1}]
    columns = ["a"]

    with tempfile.TemporaryDirectory() as tmpdir:
        path1 = format_to_excel(rows, columns, output_dir=tmpdir)
        path2 = format_to_excel(rows, columns, output_dir=tmpdir)
        assert path1 != path2
        os.unlink(path1)
        os.unlink(path2)


# ---------------------------------------------------------------------------
# execute_sql_xlsx tool tests
# ---------------------------------------------------------------------------


@pytest_asyncio.fixture
async def mock_db_connection():
    """Create a mock DB connection pool."""
    conn = MagicMock()
    conn.pool_connect = AsyncMock()
    conn.close = AsyncMock()
    return conn


@pytest.mark.asyncio
async def test_execute_sql_xlsx_success(mock_db_connection):
    """execute_sql_xlsx returns file path and row count on success."""
    mock_driver = AsyncMock()
    mock_driver.execute_query = AsyncMock(
        return_value=[
            MockRowResult({"id": 1, "name": "Alice"}),
            MockRowResult({"id": 2, "name": "Bob"}),
        ]
    )

    with (
        patch("postgres_mcp.server.current_access_mode", AccessMode.UNRESTRICTED),
        patch("postgres_mcp.server.db_connection", mock_db_connection),
        patch("postgres_mcp.server.get_sql_driver", return_value=mock_driver),
    ):
        result = await server.execute_sql_xlsx("SELECT * FROM users")

    text = response_text(result)
    assert "Excel file created:" in text
    assert "Rows exported: 2" in text
    assert "id, name" in text


@pytest.mark.asyncio
async def test_execute_sql_xlsx_empty_results(mock_db_connection):
    """execute_sql_xlsx returns informational text for empty results."""
    mock_driver = AsyncMock()
    mock_driver.execute_query = AsyncMock(return_value=[])

    with (
        patch("postgres_mcp.server.current_access_mode", AccessMode.UNRESTRICTED),
        patch("postgres_mcp.server.db_connection", mock_db_connection),
        patch("postgres_mcp.server.get_sql_driver", return_value=mock_driver),
    ):
        result = await server.execute_sql_xlsx("SELECT * FROM empty_table")

    text = response_text(result)
    assert "no results" in text.lower()
    assert not text.startswith("Error:")


@pytest.mark.asyncio
async def test_execute_sql_xlsx_none_results(mock_db_connection):
    """execute_sql_xlsx returns informational text when query returns None."""
    mock_driver = AsyncMock()
    mock_driver.execute_query = AsyncMock(return_value=None)

    with (
        patch("postgres_mcp.server.current_access_mode", AccessMode.UNRESTRICTED),
        patch("postgres_mcp.server.db_connection", mock_db_connection),
        patch("postgres_mcp.server.get_sql_driver", return_value=mock_driver),
    ):
        result = await server.execute_sql_xlsx("SELECT * FROM users")

    text = response_text(result)
    assert "no results" in text.lower()
    assert not text.startswith("Error:")


@pytest.mark.asyncio
async def test_execute_sql_xlsx_wraps_query_with_limit(mock_db_connection):
    """execute_sql_xlsx applies max_rows as an outer LIMIT."""
    mock_driver = AsyncMock()
    mock_driver.execute_query = AsyncMock(
        return_value=[
            MockRowResult({"id": 1}),
            MockRowResult({"id": 2}),
        ]
    )

    with (
        patch("postgres_mcp.server.current_access_mode", AccessMode.UNRESTRICTED),
        patch("postgres_mcp.server.db_connection", mock_db_connection),
        patch("postgres_mcp.server.get_sql_driver", return_value=mock_driver),
    ):
        await server.execute_sql_xlsx("SELECT * FROM users", max_rows=100)

    called_sql = mock_driver.execute_query.call_args[0][0]
    assert called_sql == "SELECT * FROM (\nSELECT * FROM users\n) AS _postgres_mcp_export LIMIT 100"


@pytest.mark.asyncio
async def test_execute_sql_xlsx_caps_query_with_existing_limit(mock_db_connection):
    """execute_sql_xlsx caps results even when the query already has a LIMIT."""
    mock_driver = AsyncMock()
    mock_driver.execute_query = AsyncMock(
        return_value=[
            MockRowResult({"id": 1}),
        ]
    )

    with (
        patch("postgres_mcp.server.current_access_mode", AccessMode.UNRESTRICTED),
        patch("postgres_mcp.server.db_connection", mock_db_connection),
        patch("postgres_mcp.server.get_sql_driver", return_value=mock_driver),
    ):
        await server.execute_sql_xlsx("SELECT * FROM users LIMIT 500", max_rows=100)

    called_sql = mock_driver.execute_query.call_args[0][0]
    assert called_sql == "SELECT * FROM (\nSELECT * FROM users LIMIT 500\n) AS _postgres_mcp_export LIMIT 100"


@pytest.mark.asyncio
async def test_execute_sql_xlsx_caps_query_with_nested_limit(mock_db_connection):
    """An inner LIMIT cannot bypass the outer max_rows cap."""
    mock_driver = AsyncMock()
    mock_driver.execute_query = AsyncMock(return_value=[MockRowResult({"id": 1})])

    query = "SELECT * FROM (SELECT * FROM users LIMIT 1) AS nested"
    with (
        patch("postgres_mcp.server.current_access_mode", AccessMode.UNRESTRICTED),
        patch("postgres_mcp.server.db_connection", mock_db_connection),
        patch("postgres_mcp.server.get_sql_driver", return_value=mock_driver),
    ):
        await server.execute_sql_xlsx(query, max_rows=100)

    called_sql = mock_driver.execute_query.call_args[0][0]
    assert called_sql == f"SELECT * FROM (\n{query}\n) AS _postgres_mcp_export LIMIT 100"


@pytest.mark.parametrize(
    "query",
    [
        "SELECT 'LIMIT' AS label",
        "SELECT 1 -- LIMIT in a comment",
    ],
)
@pytest.mark.asyncio
async def test_execute_sql_xlsx_caps_limit_text_in_query(query, mock_db_connection):
    """LIMIT text in a string or comment cannot bypass the outer cap."""
    mock_driver = AsyncMock()
    mock_driver.execute_query = AsyncMock(return_value=[MockRowResult({"value": 1})])

    with patch("postgres_mcp.server.get_sql_driver", return_value=mock_driver):
        await server.execute_sql_xlsx(query, max_rows=100)

    called_sql = mock_driver.execute_query.call_args[0][0]
    assert called_sql == f"SELECT * FROM (\n{query}\n) AS _postgres_mcp_export LIMIT 100"


@pytest.mark.asyncio
async def test_execute_sql_xlsx_restricted_mode_uses_safe_driver():
    """Restricted mode validates the wrapped query before database execution."""
    base_driver = AsyncMock()
    safe_driver = SafeSqlDriver(base_driver)
    write_query = "WITH deleted AS (DELETE FROM users RETURNING id) SELECT * FROM deleted"

    with patch("postgres_mcp.server.get_sql_driver", return_value=safe_driver):
        result = await server.execute_sql_xlsx(write_query)

    assert response_text(result).startswith("Error:")
    base_driver.execute_query.assert_not_awaited()


@pytest.mark.asyncio
async def test_execute_sql_xlsx_query_error(mock_db_connection):
    """execute_sql_xlsx returns error response on query failure."""
    mock_driver = AsyncMock()
    mock_driver.execute_query = AsyncMock(side_effect=Exception("Connection lost"))

    with (
        patch("postgres_mcp.server.current_access_mode", AccessMode.UNRESTRICTED),
        patch("postgres_mcp.server.db_connection", mock_db_connection),
        patch("postgres_mcp.server.get_sql_driver", return_value=mock_driver),
    ):
        result = await server.execute_sql_xlsx("SELECT * FROM users")

    text = response_text(result)
    assert text.startswith("Error:")
    assert "Connection lost" in text


@pytest.mark.asyncio
async def test_execute_sql_xlsx_max_rows_zero_rejected():
    """execute_sql_xlsx rejects max_rows=0 via validation."""
    with pytest.raises(Exception):  # noqa: B017
        await server.execute_sql_xlsx("SELECT 1", max_rows=0)


@pytest.mark.asyncio
async def test_execute_sql_xlsx_max_rows_negative_rejected():
    """execute_sql_xlsx rejects negative max_rows via validation."""
    with pytest.raises(Exception):  # noqa: B017
        await server.execute_sql_xlsx("SELECT 1", max_rows=-5)
